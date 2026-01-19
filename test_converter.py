#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
테스트용 특허 명세서 샘플 문서 생성 및 변환 테스트
"""

from docx import Document
from docx.shared import Pt
from patent_format_converter import PatentFormatConverter
import os

def create_test_patent_document():
    """테스트용 특허 명세서 Word 문서 생성"""

    doc = Document()

    # 발명의 명칭
    doc.add_paragraph('【발명의 명칭】')
    doc.add_paragraph('인공지능 기반 문서 자동 변환 시스템')
    doc.add_paragraph()

    # 기술분야
    doc.add_paragraph('【기술분야】')
    doc.add_paragraph('본 발명은 문서 처리 기술에 관한 것으로, 특히 다양한 형식의 문서를 특허 명세서 형식으로 자동 변환하는 시스템에 관한 것이다.')
    doc.add_paragraph()

    # 발명의 배경이 되는 기술
    doc.add_paragraph('【발명의 배경이 되는 기술】')
    doc.add_paragraph('종래에는 특허 명세서를 작성할 때 수동으로 복사하여 붙여넣기 하는 방식을 사용하였다.')
    doc.add_paragraph('이러한 방식은 시간이 많이 소요되고 오류가 발생하기 쉬운 문제점이 있었다.')
    doc.add_paragraph()

    # 해결하려는 과제
    doc.add_paragraph('【해결하려는 과제】')
    doc.add_paragraph('본 발명은 상기와 같은 문제점을 해결하기 위한 것으로, 문서 변환을 자동화하여 작업 효율을 높이는 것을 목적으로 한다.')
    doc.add_paragraph()

    # 과제의 해결 수단
    doc.add_paragraph('【과제의 해결 수단】')
    doc.add_paragraph('상기 목적을 달성하기 위한 본 발명은 다음과 같은 구성을 포함한다.')
    doc.add_paragraph('문서 읽기 모듈, 섹션 인식 모듈, HLT 변환 모듈을 포함하는 문서 변환 시스템.')
    doc.add_paragraph()

    # 발명의 효과
    doc.add_paragraph('【발명의 효과】')
    doc.add_paragraph('본 발명에 따르면 특허 명세서 작성 시간을 대폭 단축할 수 있다.')
    doc.add_paragraph('또한 수동 작업에서 발생하는 오류를 방지할 수 있다.')
    doc.add_paragraph()

    # 발명을 실시하기 위한 구체적인 내용
    doc.add_paragraph('【발명을 실시하기 위한 구체적인 내용】')
    doc.add_paragraph('이하, 첨부된 도면을 참조하여 본 발명의 바람직한 실시예를 상세히 설명한다.')
    doc.add_paragraph('본 시스템은 Word, HWP, PDF 파일을 입력받아 HLT 형식으로 변환한다.')
    doc.add_paragraph('변환 과정에서 문서의 구조를 자동으로 인식하고 해당하는 섹션으로 분류한다.')
    doc.add_paragraph()

    # 청구범위
    doc.add_paragraph('【청구범위】')
    doc.add_paragraph('청구항 1: 문서 파일을 읽는 문서 읽기부; 상기 문서의 섹션을 인식하는 섹션 인식부; 및 인식된 섹션을 HLT 형식으로 변환하는 변환부를 포함하는 특허 문서 변환 시스템.')
    doc.add_paragraph('청구항 2: 제1항에 있어서, 상기 문서 파일은 Word, HWP, PDF 형식 중 하나인 것을 특징으로 하는 특허 문서 변환 시스템.')
    doc.add_paragraph('청구항 3: 제1항에 있어서, 상기 변환부는 한국특허청 표준 XML 형식으로 출력하는 것을 특징으로 하는 특허 문서 변환 시스템.')
    doc.add_paragraph()

    # 요약
    doc.add_paragraph('【요약】')
    doc.add_paragraph('본 발명은 다양한 형식의 문서를 특허 명세서 형식(HLT)으로 자동 변환하는 시스템에 관한 것이다.')

    # 파일 저장
    output_path = 'test_patent.docx'
    doc.save(output_path)
    print(f"✓ 테스트 문서 생성 완료: {output_path}")
    return output_path

def test_conversion(input_file):
    """변환 테스트"""
    print()
    print("=" * 70)
    print("변환 테스트 시작")
    print("=" * 70)
    print()

    try:
        # 변환기 실행
        converter = PatentFormatConverter(input_file)
        output_file = converter.convert()

        # 결과 확인
        if os.path.exists(output_file):
            file_size = os.path.getsize(output_file)
            print()
            print("=" * 70)
            print("✓ 테스트 성공!")
            print("=" * 70)
            print()
            print(f"생성된 파일: {output_file}")
            print(f"파일 크기: {file_size} bytes")
            print()

            # HLT 파일 내용 미리보기
            print("HLT 파일 내용 미리보기:")
            print("-" * 70)
            with open(output_file, 'r', encoding='utf-8') as f:
                lines = f.readlines()
                for i, line in enumerate(lines[:30], 1):  # 처음 30줄만 출력
                    print(f"{i:3d}: {line.rstrip()}")
                if len(lines) > 30:
                    print(f"... (총 {len(lines)}줄)")
            print("-" * 70)
            print()

            return True
        else:
            print("❌ 테스트 실패: 출력 파일이 생성되지 않았습니다")
            return False

    except Exception as e:
        print(f"❌ 테스트 실패: {e}")
        import traceback
        traceback.print_exc()
        return False

def main():
    """메인 테스트 함수"""
    print("=" * 70)
    print("특허 문서 변환기 테스트")
    print("Patent Document Converter Test")
    print("=" * 70)
    print()

    # 1. 테스트 문서 생성
    input_file = create_test_patent_document()

    # 2. 변환 테스트
    success = test_conversion(input_file)

    # 3. 결과 요약
    print()
    print("=" * 70)
    if success:
        print("✓ 모든 테스트 통과!")
        print()
        print("생성된 파일:")
        print(f"  - 입력: {input_file}")
        print(f"  - 출력: test_patent.hlt")
        print()
        print("다음 단계:")
        print("  1. test_patent.hlt 파일을 K-Editor에서 열어보세요")
        print("  2. 내용이 제대로 변환되었는지 확인하세요")
    else:
        print("❌ 테스트 실패")
    print("=" * 70)

    return 0 if success else 1

if __name__ == '__main__':
    import sys
    sys.exit(main())
